"""Build a fictional gerontology manuscript with common formatting problems."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.0
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.75); s.top_margin = s.bottom_margin = Inches(0.75)

def para(text, bold=False, size=None, align=None, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

para("Social Participation and Depressive Symptoms Among Older Adults: Evidence From a Fictional Panel Study", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para("DEMONSTRATION MANUSCRIPT. All authors, data, and results are invented to show the skill's formatting tools.", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("A. Example, PhD, and B. Sample, MS", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Abstract", bold=True)
para("Social participation is widely thought to protect mental health in later life, yet much of the evidence comes from cross-sectional data, and less is known about whether the benefits differ by gender or by the type of activity. In this study we used four waves of a fictional panel study of adults aged 65 and older (N = 2,431) to examine the association between social participation and depressive symptoms over six years. Social participation was measured as the number of activities in which respondents took part at least monthly, including volunteering, religious services, clubs, and visits with friends. Depressive symptoms were measured with an eight-item scale. We estimated linear mixed models with random intercepts and adjusted for age, gender, education, wealth, marital status, chronic conditions, and functional limitations. Each additional activity was associated with fewer depressive symptoms at baseline and with a slower increase in symptoms over time. The association was stronger for volunteering and visits with friends than for religious attendance, and it was somewhat stronger among men than among women, although the gender difference was modest. Results were robust to the exclusion of respondents with probable depression at baseline and to alternative codings of participation. These findings suggest that sustaining opportunities for social participation, particularly volunteering and informal social contact, may help slow the rise of depressive symptoms in later life, and that programs should consider the specific activities that older men and women prefer. Future research should examine the mechanisms that link participation to mental health, including social support, sense of purpose, and physical activity.")
para("Keywords: social participation; depression; aging; volunteering; longitudinal")

para("INTRODUCTION", bold=True)
intro = ["Depressive symptoms are common in later life and are associated with disability, health care use, and mortality. Social participation, meaning involvement in activities that bring people into contact with others, has been proposed as a modifiable factor that may protect mental health as people age. Activity theory and the socioemotional selectivity framework both suggest that meaningful engagement supports well-being, although they differ in which activities they expect to matter most.",
         "Prior studies have generally found that older adults who participate in more activities report fewer depressive symptoms. However, many of these studies are cross-sectional, so they cannot rule out the possibility that people with better mental health are simply more likely to participate. Longitudinal studies have produced more mixed results, and few have compared types of activities or examined differences by gender.",
         "In this study, we addressed these gaps using four waves of panel data. We asked three questions. First, is social participation associated with the level and the trajectory of depressive symptoms? Second, do the associations differ by type of activity? Third, do they differ between older men and women?"]
for t in intro: para(t)
para("METHODS", bold=True)
for t in ["Data came from a fictional panel study of community-dwelling adults aged 65 and older, interviewed every two years from 2014 to 2020. The analytic sample included 2,431 respondents who completed the baseline interview and at least one follow-up interview.",
          "Depressive symptoms were measured with an eight-item scale (range 0-8). Social participation was the number of activities, out of six, in which respondents took part at least once a month. Covariates included age, gender, education, household wealth, marital status, number of chronic conditions, and functional limitations.",
          "We estimated linear mixed models with random intercepts for respondents. Model 1 included participation and time; Model 2 added the covariates; Model 3 added an interaction between participation and time. All analyses were conducted in Stata 18."]:
    para(t)
para("RESULTS", bold=True)
para("Table 1 shows the characteristics of the sample at baseline. Table 2 presents the results of the mixed models. Each additional activity was associated with 0.21 fewer depressive symptoms (p = 0.000), and the association remained after adjustment for covariates.")

# ---- Table 1: descriptive, with typical problems
para("Table 1. Sample characteristics at baseline (N=2431)", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
rows = [["Characteristic", "Men (n=1032)", "Women (n=1399)", "p-value"],
        ["Age, mean (SD)", "74.2 (6.81)", "75.31 (7.2)", "0.000"],
        ["Education, years, mean (SD)", "12.4 (3.1)", "11.9 (2.95)", "0.002"],
        ["Married, n (%)", "721 (69.9)", "655 (46.82)", "<0.001"],
        ["Chronic conditions, mean (SD)", "2.1 (1.4)", "2.3 (1.38)", "0.01"],
        ["Depressive symptoms, mean (SD)", "1.6 (1.9)", "2.04 (2.1)", "0.000"],
        ["Activities per month, mean (SD)", "2.8 (1.6)", "3.1 (1.7)", "0.000"]]
t = doc.add_table(rows=len(rows), cols=4); t.style = "Table Grid"
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = t.cell(i, j); c.text = v
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        if i == 0:
            c.paragraphs[0].runs[0].bold = True
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3")
            c._tc.get_or_add_tcPr().append(sh)
para("Note: SD = standard deviation. p-values from t-tests and chi-square tests.", size=9, align=WD_ALIGN_PARAGRAPH.LEFT)

# ---- Table 2: regression
para("Table 2. Linear mixed models of depressive symptoms", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
rows = [["", "Model 1", "Model 2", "Model 3"],
        ["Activities (number)", "-0.31***", "-0.21***", "-0.19***"],
        ["", "(0.03)", "(0.034)", "(0.03)"],
        ["Time (waves)", "0.12***", "0.1***", "0.15***"],
        ["", "(0.02)", "(0.02)", "(0.021)"],
        ["Activities x Time", "", "", "-0.02*"],
        ["", "", "", "(0.01)"],
        ["Women", "", "0.274**", "0.27**"],
        ["", "", "(0.09)", "(0.09)"],
        ["Observations", "8,904", "8904", "8,904"],
        ["Respondents", "2431", "2,431", "2,431"]]
t = doc.add_table(rows=len(rows), cols=4); t.style = "Table Grid"
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = t.cell(i, j); c.text = v
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        if i == 0 and v: c.paragraphs[0].runs[0].bold = True
        if j == 3 and i == 5: c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
para("Note: Standard errors in parentheses. Model 2 adjusts for age, education, wealth, marital status, chronic conditions, and functional limitations. * p<0.05, ** p<0.01, *** p<0.001", size=9, align=WD_ALIGN_PARAGRAPH.LEFT)

para("DISCUSSION", bold=True)
para("In this fictional sample, older adults who took part in more social activities reported fewer depressive symptoms, and their symptoms rose more slowly over time. These patterns are consistent with prior longitudinal work and suggest that programs supporting volunteering and informal social contact may benefit mental health in later life. Because the data are observational, the associations should not be read as causal effects.")
para("REFERENCES", bold=True)
for r in ["Example, A. (2020). Social participation in later life. Journal of Fictional Aging, 12(3), 45-67.",
          "Sample, B., & Example, A. (2022). Depressive symptoms and activity: A panel study. Fictional Gerontology Review, 8, 101-120."]:
    para(r, align=WD_ALIGN_PARAGRAPH.LEFT)
doc.save(sys.argv[1])
